import os
import re
import json
from pathlib import Path
import openpyxl

# ============================================================
# EXCEL SKILL 1 - AUTO CHECKER
# 20 Questions | 10 Marks | 100+ Students
#
# IMPORTANT:
# Range-based questions accept BOTH:
#   1) Range WITHOUT header
#   2) Range WITH header
#
# Example:
#   A4:C13
#   A3:C13
#
# Absolute references are also accepted:
#   $A$4:$C$13
#   $A$3:$C$13
# ============================================================


BASE_DIR = Path(__file__).parent
SUBMISSIONS_DIR = BASE_DIR / "submissions"
RESULTS_DIR = BASE_DIR / "results"

QUESTION_MARK = 0.5


# ============================================================
# ANSWER CELLS
# ============================================================

ANSWER_CELLS = {
    "Q1": ("VLOOKUP", "C19"),
    "Q2": ("VLOOKUP", "C20"),
    "Q3": ("VLOOKUP", "C21"),
    "Q4": ("VLOOKUP", "C22"),

    "Q5": ("SUMIF & COUNTIF", "C20"),
    "Q6": ("SUMIF & COUNTIF", "C21"),
    "Q7": ("SUMIF & COUNTIF", "C22"),
    "Q8": ("SUMIF & COUNTIF", "C23"),
    "Q9": ("SUMIF & COUNTIF", "C24"),
    "Q10": ("SUMIF & COUNTIF", "C25"),

    "Q11": ("LEFT RIGHT MID", "C15"),
    "Q12": ("LEFT RIGHT MID", "C16"),
    "Q13": ("LEFT RIGHT MID", "C17"),
    "Q14": ("LEFT RIGHT MID", "C18"),
    "Q15": ("LEFT RIGHT MID", "C19"),
    "Q16": ("LEFT RIGHT MID", "C20"),

    "Q17": ("Complex Challenge", "C17"),
    "Q18": ("Complex Challenge", "C18"),
    "Q19": ("Complex Challenge", "C19"),
    "Q20": ("Complex Challenge", "C20"),
}


# ============================================================
# HELPERS
# ============================================================

def normalize_formula(value):
    """
    Normalizes formula for comparison.

    Examples:
        =VLOOKUP(...)     -> VLOOKUP(...)
        $A$4:$C$13       -> A4:C13
        a4:c13           -> A4:C13
        spaces removed
    """

    if value is None:
        return ""

    s = str(value).strip().upper()

    # Remove spaces
    s = re.sub(r"\s+", "", s)

    # Remove leading =
    if s.startswith("="):
        s = s[1:]

    # Remove absolute reference $
    s = s.replace("$", "")

    return s


def normalize_range_list(ranges):
    """
    Converts possible ranges into normalized set.

    Example:
        ["A4:C13", "A3:C13"]
    becomes:
        {"A4:C13", "A3:C13"}
    """

    return {
        normalize_formula(r)
        for r in ranges
    }


def is_formula(value):
    return (
        isinstance(value, str)
        and value.strip().startswith("=")
    )


def get_function(formula):
    f = normalize_formula(formula)

    m = re.match(
        r"^([A-Z][A-Z0-9_.]*)\(",
        f
    )

    return m.group(1) if m else ""


def split_args(formula):
    """
    Safely split Excel function arguments.

    Handles nested functions such as:

    =LEFT(C4,FIND("@",C4)-1)
    """

    f = normalize_formula(formula)

    p = f.find("(")
    q = f.rfind(")")

    if p < 0 or q < 0:
        return []

    inside = f[p + 1:q]

    args = []
    current = []
    depth = 0
    quote = False

    for ch in inside:

        if ch == '"':
            quote = not quote
            current.append(ch)

        elif not quote and ch == "(":
            depth += 1
            current.append(ch)

        elif not quote and ch == ")":
            depth -= 1
            current.append(ch)

        elif not quote and ch == "," and depth == 0:
            args.append(
                "".join(current).strip()
            )
            current = []

        else:
            current.append(ch)

    if current:
        args.append(
            "".join(current).strip()
        )

    return args


# ============================================================
# VLOOKUP CHECKER
# ============================================================

def check_vlookup(
    formula,
    lookup_values,
    return_col,
    table_ranges
):

    result = {
        "marks": 0,
        "status": "WRONG",
        "issues": [],
        "checks": {}
    }

    # --------------------------------------------------------
    # Formula check
    # --------------------------------------------------------

    if not is_formula(formula):
        result["issues"].append(
            "No formula found."
        )
        return result

    # --------------------------------------------------------
    # Function check
    # --------------------------------------------------------

    if get_function(formula) != "VLOOKUP":
        result["issues"].append(
            "VLOOKUP is required."
        )
        return result

    # --------------------------------------------------------
    # Arguments
    # --------------------------------------------------------

    args = split_args(formula)

    if len(args) < 4:
        result["issues"].append(
            "VLOOKUP needs 4 arguments."
        )
        return result

    lookup = normalize_formula(args[0])
    table = normalize_formula(args[1])
    col = normalize_formula(args[2])
    match = normalize_formula(args[3])

    # --------------------------------------------------------
    # Possible lookup values
    #
    # Example Q2:
    # "Sara Khan"
    # "B20"
    # Sara Khan
    # B20
    # --------------------------------------------------------

    possible_lookup_values = set()

    for value in lookup_values:

        value = normalize_formula(value)

        possible_lookup_values.add(value)

        # Also accept quoted text
        possible_lookup_values.add(
            f'"{value}"'
        )

    # --------------------------------------------------------
    # Possible table ranges
    #
    # Example:
    #
    # B4:E13  -> WITHOUT HEADER
    # B3:E13  -> WITH HEADER
    #
    # Both are valid.
    # --------------------------------------------------------

    possible_table_ranges = normalize_range_list(
        table_ranges
    )

    # --------------------------------------------------------
    # Checks
    # --------------------------------------------------------

    result["checks"]["function"] = 1

    result["checks"]["lookup_value"] = int(
        lookup in possible_lookup_values
    )

    result["checks"]["table_range"] = int(
        table in possible_table_ranges
    )

    result["checks"]["return_column"] = int(
        col == str(return_col)
    )

    result["checks"]["exact_match"] = int(
        match in {"0", "FALSE"}
    )

    # --------------------------------------------------------
    # Issues
    # --------------------------------------------------------

    for name, value in result["checks"].items():

        if value == 0:
            result["issues"].append(
                name + " is incorrect."
            )

    # --------------------------------------------------------
    # Final
    # --------------------------------------------------------

    if all(
        value == 1
        for value in result["checks"].values()
    ):
        result["marks"] = QUESTION_MARK
        result["status"] = "CORRECT"

    return result


# ============================================================
# SUMIF CHECKER
# ============================================================

def check_sumif(
    formula,
    criteria_ranges,
    criteria,
    sum_ranges
):

    result = {
        "marks": 0,
        "status": "WRONG",
        "issues": [],
        "checks": {}
    }

    if not is_formula(formula):
        result["issues"].append(
            "No formula found."
        )
        return result

    if get_function(formula) != "SUMIF":
        result["issues"].append(
            "SUMIF is required."
        )
        return result

    args = split_args(formula)

    if len(args) != 3:
        result["issues"].append(
            "SUMIF must have 3 arguments."
        )
        return result

    criteria_range = normalize_formula(
        args[0]
    )

    actual_criteria = normalize_formula(
        args[1]
    )

    sum_range = normalize_formula(
        args[2]
    )

    # --------------------------------------------------------
    # Possible criteria ranges
    # --------------------------------------------------------

    possible_criteria_ranges = (
        normalize_range_list(
            criteria_ranges
        )
    )

    # --------------------------------------------------------
    # Possible sum ranges
    # --------------------------------------------------------

    possible_sum_ranges = (
        normalize_range_list(
            sum_ranges
        )
    )

    # --------------------------------------------------------
    # Possible criteria
    # --------------------------------------------------------

    possible_criteria = {
        normalize_formula(criteria),
        f'"{normalize_formula(criteria)}"'
    }

    # --------------------------------------------------------
    # Checks
    # --------------------------------------------------------

    result["checks"]["function"] = 1

    result["checks"]["criteria_range"] = int(
        criteria_range in possible_criteria_ranges
    )

    result["checks"]["criteria"] = int(
        actual_criteria in possible_criteria
    )

    result["checks"]["sum_range"] = int(
        sum_range in possible_sum_ranges
    )

    # --------------------------------------------------------
    # Issues
    # --------------------------------------------------------

    for name, value in result["checks"].items():

        if value == 0:
            result["issues"].append(
                name + " is incorrect."
            )

    # --------------------------------------------------------
    # Final
    # --------------------------------------------------------

    if all(
        value == 1
        for value in result["checks"].values()
    ):
        result["marks"] = QUESTION_MARK
        result["status"] = "CORRECT"

    return result


# ============================================================
# COUNTIF CHECKER
# ============================================================

def check_countif(
    formula,
    criteria_ranges,
    criteria
):

    result = {
        "marks": 0,
        "status": "WRONG",
        "issues": [],
        "checks": {}
    }

    if not is_formula(formula):
        result["issues"].append(
            "No formula found."
        )
        return result

    if get_function(formula) != "COUNTIF":
        result["issues"].append(
            "COUNTIF is required."
        )
        return result

    args = split_args(formula)

    if len(args) != 2:
        result["issues"].append(
            "COUNTIF must have 2 arguments."
        )
        return result

    criteria_range = normalize_formula(
        args[0]
    )

    actual_criteria = normalize_formula(
        args[1]
    )

    # --------------------------------------------------------
    # Possible ranges
    # --------------------------------------------------------

    possible_criteria_ranges = (
        normalize_range_list(
            criteria_ranges
        )
    )

    # --------------------------------------------------------
    # Possible criteria
    # --------------------------------------------------------

    possible_criteria = {
        normalize_formula(criteria),
        f'"{normalize_formula(criteria)}"'
    }

    # --------------------------------------------------------
    # Checks
    # --------------------------------------------------------

    result["checks"]["function"] = 1

    result["checks"]["criteria_range"] = int(
        criteria_range in possible_criteria_ranges
    )

    result["checks"]["criteria"] = int(
        actual_criteria in possible_criteria
    )

    # --------------------------------------------------------
    # Issues
    # --------------------------------------------------------

    for name, value in result["checks"].items():

        if value == 0:
            result["issues"].append(
                name + " is incorrect."
            )

    # --------------------------------------------------------
    # Final
    # --------------------------------------------------------

    if all(
        value == 1
        for value in result["checks"].values()
    ):
        result["marks"] = QUESTION_MARK
        result["status"] = "CORRECT"

    return result


# ============================================================
# TEXT FORMULA CHECKER
# ============================================================

def check_text_formula(
    formula,
    accepted_patterns,
    expected_text
):

    result = {
        "marks": 0,
        "status": "WRONG",
        "issues": [],
        "checks": {}
    }

    if not is_formula(formula):
        result["issues"].append(
            "No formula found."
        )
        return result

    f = normalize_formula(formula)

    result["checks"]["formula_structure"] = int(
        any(
            re.fullmatch(pattern, f)
            for pattern in accepted_patterns
        )
    )

    if result["checks"]["formula_structure"]:

        result["marks"] = QUESTION_MARK
        result["status"] = "CORRECT"

    else:

        result["issues"].append(
            "Formula does not match an accepted "
            f"solution for expected result: {expected_text}"
        )

    return result


# ============================================================
# Q1-Q4 VLOOKUP
#
# BOTH ranges are accepted:
#
# WITHOUT HEADER
# WITH HEADER
# ============================================================

def check_q1(f):

    return check_vlookup(
        f,
        ["E003"],
        3,
        [
            "A4:C13",   # WITHOUT HEADER
            "A3:C13"    # WITH HEADER
        ]
    )


def check_q2(f):

    return check_vlookup(
        f,

        # Both direct value and cell reference accepted
        [
            "Sara Khan",
            "B20"
        ],

        4,

        [
            "B4:E13",   # WITHOUT HEADER
            "B3:E13"    # WITH HEADER
        ]
    )


def check_q3(f):

    return check_vlookup(
        f,
        ["E007"],
        4,
        [
            "A4:D13",   # WITHOUT HEADER
            "A3:D13"    # WITH HEADER
        ]
    )


def check_q4(f):

    return check_vlookup(
        f,
        ["E010"],
        2,
        [
            "A4:B13",   # WITHOUT HEADER
            "A3:B13"    # WITH HEADER
        ]
    )


# ============================================================
# Q5-Q10 SUMIF / COUNTIF
#
# BOTH WITH HEADER AND WITHOUT HEADER ACCEPTED
# ============================================================

def check_q5(f):

    return check_sumif(
        f,

        [
            "B4:B13",   # WITHOUT HEADER
            "B3:B13"    # WITH HEADER
        ],

        "ALI",

        [
            "F4:F13",   # WITHOUT HEADER
            "F3:F13"    # WITH HEADER
        ]
    )


def check_q6(f):

    return check_countif(
        f,

        [
            "C4:C13",   # WITHOUT HEADER
            "C3:C13"    # WITH HEADER
        ],

        "LAPTOP"
    )


def check_q7(f):

    return check_sumif(
        f,

        [
            "B4:B13",   # WITHOUT HEADER
            "B3:B13"    # WITH HEADER
        ],

        "SARA",

        [
            "E4:E13",   # WITHOUT HEADER
            "E3:E13"    # WITH HEADER
        ]
    )


def check_q8(f):

    return check_countif(
        f,

        [
            "D4:D13",   # WITHOUT HEADER
            "D3:D13"    # WITH HEADER
        ],

        "ELECTRONICS"
    )


def check_q9(f):

    return check_sumif(
        f,

        [
            "D4:D13",   # WITHOUT HEADER
            "D3:D13"    # WITH HEADER
        ],

        "ACCESSORIES",

        [
            "F4:F13",   # WITHOUT HEADER
            "F3:F13"    # WITH HEADER
        ]
    )


def check_q10(f):

    return check_countif(
        f,

        [
            "B4:B13",   # WITHOUT HEADER
            "B3:B13"    # WITH HEADER
        ],

        "FATIMA"
    )


# ============================================================
# Q11-Q16 TEXT FUNCTIONS
# ============================================================

def check_q11(f):

    patterns = [
        r'LEFT\(A4,3\)',
    ]

    return check_text_formula(
        f,
        patterns,
        "Ahm"
    )


def check_q12(f):

    patterns = [
        r'RIGHT\(B4,7\)',
    ]

    return check_text_formula(
        f,
        patterns,
        "1234567"
    )


def check_q13(f):

    # Username from:
    # ahmed.ali@gmail.com
    #
    # Result:
    # ahmed.ali

    patterns = [
        r'LEFT\(C4,FIND\("@",C4\)-1\)',
    ]

    return check_text_formula(
        f,
        patterns,
        "ahmed.ali"
    )


def check_q14(f):

    # STU-2024-001
    # Result:
    # 2024

    patterns = [
        r'MID\(D4,5,4\)',
    ]

    return check_text_formula(
        f,
        patterns,
        "2024"
    )


def check_q15(f):

    # sara.f@hotmail.com
    # Result:
    # hotmail.com

    patterns = [
        r'RIGHT\(C5,LEN\(C5\)-FIND\("@",C5\)\)',

        r'MID\(C5,FIND\("@",C5\)+1,LEN\(C5\)\)',
    ]

    return check_text_formula(
        f,
        patterns,
        "hotmail.com"
    )


def check_q16(f):

    # Sara Fatima
    # Result:
    # Fatima

    patterns = [
        r'RIGHT\(A5,LEN\(A5\)-FIND\(" ",A5\)\)',

        r'MID\(A5,FIND\(" ",A5\)+1,LEN\(A5\)\)',
    ]

    return check_text_formula(
        f,
        patterns,
        "Fatima"
    )


# ============================================================
# Q17-Q20 COMPLEX CHALLENGE
# ============================================================

def check_q17(f):

    patterns = [
        r'D4\*E4',
        r'E4\*D4',
    ]

    return check_text_formula(
        f,
        patterns,
        "1020000"
    )


def check_q18(f):

    patterns = [
        r'RIGHT\(A4,3\)',
        r'MID\(A4,9,3\)',
    ]

    return check_text_formula(
        f,
        patterns,
        "LAP"
    )


def check_q19(f):

    return check_countif(
        f,

        [
            "C4:C11",   # WITHOUT HEADER
            "C3:C11"    # WITH HEADER
        ],

        "ACCESSORIES"
    )


def check_q20(f):

    return check_sumif(
        f,

        [
            "C4:C11",   # WITHOUT HEADER
            "C3:C11"    # WITH HEADER
        ],

        "ELECTRONICS",

        [
            "F4:F11",   # WITHOUT HEADER
            "F3:F11"    # WITH HEADER
        ]
    )


# ============================================================
# CHECKER MAP
# ============================================================

CHECKERS = {

    "Q1": check_q1,
    "Q2": check_q2,
    "Q3": check_q3,
    "Q4": check_q4,

    "Q5": check_q5,
    "Q6": check_q6,
    "Q7": check_q7,
    "Q8": check_q8,
    "Q9": check_q9,
    "Q10": check_q10,

    "Q11": check_q11,
    "Q12": check_q12,
    "Q13": check_q13,
    "Q14": check_q14,
    "Q15": check_q15,
    "Q16": check_q16,

    "Q17": check_q17,
    "Q18": check_q18,
    "Q19": check_q19,
    "Q20": check_q20,
}


# ============================================================
# IF & NESTED IF DIAGNOSTIC
# ============================================================

def check_grade_formula(formula, row):

    if not is_formula(formula):
        return 0, "No formula"

    f = normalize_formula(formula)

    required = [
        f"C{row}",
        "IF(",
        '"A+"',
        '"A"',
        '"B"',
        '"C"',
        '"D"',
        '"F"',
    ]

    ok = all(
        x.upper() in f
        for x in required
    )

    if ok:
        return 1, "Valid nested IF grade formula"

    return (
        0,
        "Grade formula missing required "
        "IF/threshold/grade logic"
    )


def check_status_formula(formula, row):

    if not is_formula(formula):
        return 0, "No formula"

    f = normalize_formula(formula)

    required = [
        f"C{row}",
        "IF(",
        '"PASS"',
        '"FAIL"',
    ]

    ok = all(
        x.upper() in f
        for x in required
    )

    if ok:
        return 1, "Valid IF status formula"

    return (
        0,
        "Status formula missing required "
        "IF/Pass/Fail logic"
    )


def check_if_section(ws):

    rows = []

    for row in range(4, 14):

        grade_formula = ws[f"D{row}"].value
        status_formula = ws[f"E{row}"].value

        grade_ok, grade_msg = check_grade_formula(
            grade_formula,
            row
        )

        status_ok, status_msg = check_status_formula(
            status_formula,
            row
        )

        rows.append({

            "Row": row,

            "Grade_Cell": f"D{row}",
            "Grade_0_1": grade_ok,
            "Grade_Status": grade_msg,

            "Status_Cell": f"E{row}",
            "Status_0_1": status_ok,
            "Status_Status": status_msg,
        })

    return rows


# ============================================================
# CHECK ONE STUDENT FILE
# ============================================================

def check_student_file(file_path):

    result = {
        "Student": file_path.stem,
        "Total": 0,
        "Percentage": 0,
        "Questions": {},
        "IF_Diagnostic": [],
        "Error": ""
    }

    try:

        keep_vba = (
            file_path.suffix.lower()
            == ".xlsm"
        )

        wb = openpyxl.load_workbook(
            file_path,
            data_only=False,
            keep_vba=keep_vba
        )

        total = 0

        # ----------------------------------------------------
        # Q1-Q20
        # ----------------------------------------------------

        for question, (
            sheet_name,
            cell
        ) in ANSWER_CELLS.items():

            if sheet_name not in wb.sheetnames:

                check = {
                    "marks": 0,
                    "status": "WRONG",
                    "formula": None,
                    "issues": [
                        f"Sheet '{sheet_name}' not found."
                    ],
                    "checks": {}
                }

            else:

                ws = wb[sheet_name]

                formula = ws[cell].value

                try:

                    check = CHECKERS[
                        question
                    ](formula)

                except Exception as e:

                    check = {
                        "marks": 0,
                        "status": "ERROR",
                        "formula": formula,
                        "issues": [str(e)],
                        "checks": {}
                    }

                check["formula"] = formula

            result["Questions"][
                question
            ] = check

            total += check["marks"]

        # ----------------------------------------------------
        # IF Diagnostic
        # ----------------------------------------------------

        if "IF & NESTED IF" in wb.sheetnames:

            result["IF_Diagnostic"] = (
                check_if_section(
                    wb["IF & NESTED IF"]
                )
            )

        # ----------------------------------------------------
        # Total
        # ----------------------------------------------------

        result["Total"] = round(
            total,
            2
        )

        result["Percentage"] = round(
            (total / 10) * 100,
            2
        )

    except Exception as e:

        result["Error"] = str(e)

    return result


# ============================================================
# CREATE EXCEL REPORT
# ============================================================

def create_report(results):

    RESULTS_DIR.mkdir(
        exist_ok=True
    )

    output = (
        RESULTS_DIR
        / "Excel_Skill_1_Marks_Report.xlsx"
    )

    wb = openpyxl.Workbook()

    # ========================================================
    # MARKS SHEET
    # ========================================================

    ws = wb.active

    ws.title = "Marks"

    headers = (
        ["Student"]
        + [f"Q{i}" for i in range(1, 21)]
        + [
            "Total / 10",
            "Percentage",
            "Error"
        ]
    )

    ws.append(headers)

    for result in results:

        row = [
            result["Student"]
        ]

        for i in range(1, 21):

            q = result[
                "Questions"
            ].get(
                f"Q{i}",
                {}
            )

            row.append(
                q.get(
                    "marks",
                    0
                )
            )

        row.extend([
            result["Total"],
            result["Percentage"],
            result["Error"]
        ])

        ws.append(row)

    # ========================================================
    # DETAILED CHECKING
    # ========================================================

    detail = wb.create_sheet(
        "Detailed Checking"
    )

    detail.append([
        "Student",
        "Question",
        "Sheet",
        "Answer Cell",
        "Formula",
        "0_1",
        "Marks",
        "Status",
        "Issues"
    ])

    for result in results:

        for question, (
            sheet_name,
            cell
        ) in ANSWER_CELLS.items():

            q = result[
                "Questions"
            ].get(
                question,
                {}
            )

            detail.append([

                result["Student"],

                question,

                sheet_name,

                cell,

                q.get("formula"),

                1 if q.get(
                    "status"
                ) == "CORRECT" else 0,

                q.get(
                    "marks",
                    0
                ),

                q.get(
                    "status"
                ),

                " | ".join(
                    q.get(
                        "issues",
                        []
                    )
                )
            ])

    # ========================================================
    # FORMULA ANALYSIS
    # ========================================================

    analysis = wb.create_sheet(
        "Formula Analysis"
    )

    analysis.append([
        "Student",
        "Question",
        "Formula",
        "Status",
        "Marks",
        "Checks",
        "Issues"
    ])

    for result in results:

        for question in [
            f"Q{i}"
            for i in range(1, 21)
        ]:

            q = result[
                "Questions"
            ].get(
                question,
                {}
            )

            analysis.append([

                result["Student"],

                question,

                q.get("formula"),

                q.get("status"),

                q.get(
                    "marks",
                    0
                ),

                json.dumps(
                    q.get(
                        "checks",
                        {}
                    )
                ),

                " | ".join(
                    q.get(
                        "issues",
                        []
                    )
                )
            ])

    # ========================================================
    # IF DIAGNOSTIC
    # ========================================================

    if_sheet = wb.create_sheet(
        "IF Diagnostic"
    )

    if_sheet.append([
        "Student",
        "Row",
        "Grade Cell",
        "Grade 0/1",
        "Grade Status",
        "Status Cell",
        "Status 0/1",
        "Status Status"
    ])

    for result in results:

        for item in result[
            "IF_Diagnostic"
        ]:

            if_sheet.append([

                result["Student"],

                item["Row"],

                item["Grade_Cell"],

                item["Grade_0_1"],

                item["Grade_Status"],

                item["Status_Cell"],

                item["Status_0_1"],

                item["Status_Status"],
            ])

    # ========================================================
    # FORMATTING
    # ========================================================

    for sheet in wb.worksheets:

        sheet.freeze_panes = "A2"

        for column in sheet.columns:

            max_length = 0

            letter = (
                column[0]
                .column_letter
            )

            for cell in column:

                try:

                    max_length = max(
                        max_length,
                        len(
                            str(
                                cell.value
                            )
                        )
                    )

                except Exception:
                    pass

            sheet.column_dimensions[
                letter
            ].width = min(
                max_length + 2,
                60
            )

    wb.save(output)

    return output


# ============================================================
# MAIN
# ============================================================

def main():

    print("=" * 65)

    print(
        "       EXCEL SKILL 1 - AUTO CHECKER"
    )

    print("=" * 65)

    SUBMISSIONS_DIR.mkdir(
        exist_ok=True
    )

    RESULTS_DIR.mkdir(
        exist_ok=True
    )

    files = []

    for extension in [
        "*.xlsx",
        "*.xlsm"
    ]:

        files.extend(
            SUBMISSIONS_DIR.glob(
                extension
            )
        )

    if not files:

        print(
            "\nNo student Excel files found."
        )

        print(
            f"Put all student files inside:\n"
            f"{SUBMISSIONS_DIR}"
        )

        return

    files = sorted(files)

    print(
        f"\nFound {len(files)} student files.\n"
    )

    results = []

    for index, file_path in enumerate(
        files,
        start=1
    ):

        print(
            f"[{index}/{len(files)}] "
            f"Checking {file_path.name} ..."
        )

        result = check_student_file(
            file_path
        )

        results.append(
            result
        )

        print(
            f"    Score: "
            f"{result['Total']}/10 "
            f"({result['Percentage']}%)"
        )

    report = create_report(
        results
    )

    print(
        "\n" + "=" * 65
    )

    print(
        "CHECKING COMPLETE"
    )

    print(
        "=" * 65
    )

    print(
        f"Students checked : "
        f"{len(results)}"
    )

    print(
        f"Report           : "
        f"{report}"
    )


# ============================================================
# RUN
# ============================================================

if __name__ == "__main__":
    main()
